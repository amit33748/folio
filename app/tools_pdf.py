"""PDF and Markdown tools."""
from __future__ import annotations

import difflib
import html
import io
import re
import shutil
from pathlib import Path

import pikepdf
import pymupdf as fitz

from .registry import (
    num,
    IMAGES, OFFICE_PPT, OFFICE_WORD, OFFICE_XLS, PAGES_OPT, PDF, Result, ToolError,
    check_public_url, shrink_office_media, for_each, ghostscript, hex_color, open_pdf, out_path, page_set,
    parse_ranges, pike, run, single_or_zip, soffice, tool, zip_files,
)


# --------------------------------------------------------------------------- #
# Markdown (MarkItDown)
# --------------------------------------------------------------------------- #

_md = None


def markitdown():
    global _md
    if _md is None:
        from markitdown import MarkItDown

        _md = MarkItDown(enable_plugins=False)
    return _md


MD_ANY = sorted(
    set(PDF + OFFICE_WORD + OFFICE_PPT + OFFICE_XLS + IMAGES)
    | {".html", ".htm", ".json", ".xml", ".epub", ".zip", ".msg", ".md", ".ipynb", ".mp3", ".wav", ".m4a"}
)


def _to_markdown(files: list[Path], opts: dict, work: Path) -> Result:
    docs = []
    for f in files:
        try:
            res = markitdown().convert(str(f))
        except Exception as e:  # noqa: BLE001 - surface converter errors to the user
            raise ToolError(f"Could not convert {f.name}: {e}") from e
        docs.append({"name": f.stem + ".md", "markdown": res.markdown or "", "title": res.title or f.stem})
    return Result("markdown", docs=docs)


tool(
    id="any-to-markdown", name="Anything to Markdown", category="markdown", glyph="md", popular=True,
    keywords="llm convert extract text docx pptx xlsx epub",
    description="PDF, Word, PowerPoint, Excel, HTML, EPUB, ZIP, JSON, images — all to Markdown.",
    accept=MD_ANY, multiple=True,
)(_to_markdown)
tool(
    id="pdf-to-markdown", name="PDF to Markdown", category="markdown", glyph="md",
    description="Extract headings, lists and text from PDFs into structured Markdown.",
    accept=PDF, multiple=True,
)(_to_markdown)
tool(
    id="office-to-markdown", name="Office to Markdown", category="markdown", glyph="md",
    description="Word, PowerPoint and Excel to Markdown, with tables preserved.",
    accept=OFFICE_WORD + OFFICE_PPT + OFFICE_XLS, multiple=True,
)(_to_markdown)


@tool(
    id="url-to-markdown", name="Web page to Markdown", category="markdown", glyph="globe",
    description="Paste a URL — articles, Wikipedia, YouTube transcripts — get Markdown back.",
    accept=[], min_files=0,
    options=[{"name": "url", "label": "URL", "type": "url", "placeholder": "https://", "required": True, "default": ""}],
)
def url_to_markdown(files, opts, work):
    url = check_public_url(opts.get("url"))
    try:
        res = markitdown().convert(url)
    except Exception as e:  # noqa: BLE001
        raise ToolError(f"Could not fetch {url}: {e}") from e
    name = re.sub(r"[^\w.-]+", "_", url.split("://", 1)[1])[:60].strip("_") or "page"
    return Result("markdown", docs=[{"name": name + ".md", "markdown": res.markdown or "", "title": res.title or url}])


# --------------------------------------------------------------------------- #
# Organize
# --------------------------------------------------------------------------- #


@tool(
    id="merge", name="Merge PDF", category="organize", glyph="merge", popular=True, keywords="combine join",
    description="Combine several PDFs into one, in the order you choose.",
    accept=PDF, multiple=True, min_files=2, reorder=True,
)
def merge(files, opts, work):
    out = work / "out" / "merged.pdf"
    with pikepdf.new() as pdf:
        for f in files:
            pdf.pages.extend(pike(f).pages)
        pdf.save(out)
    return Result("file", out)


@tool(
    id="split", name="Split PDF", category="organize", glyph="split",
    description="Break a PDF into ranges or into single pages.",
    accept=PDF,
    options=[
        {"name": "mode", "label": "Split mode", "type": "segmented", "default": "ranges",
         "choices": [["ranges", "By ranges"], ["every", "Every page"], ["chunks", "Fixed size"]]},
        {"name": "ranges", "label": "Ranges", "type": "text", "placeholder": "1-3, 4-6, 7-",
         "help": "Each range becomes its own file.", "default": "", "when": {"mode": "ranges"}},
        {"name": "size", "label": "Pages per file", "type": "number", "min": 1, "default": 2, "when": {"mode": "chunks"}},
    ],
)
def split(files, opts, work):
    src = pike(files[0])
    n = len(src.pages)
    mode = opts.get("mode", "ranges")
    if mode == "every":
        ranges = [(i, i) for i in range(1, n + 1)]
    elif mode == "chunks":
        k = max(1, int(opts.get("size") or 1))
        ranges = [(i, min(i + k - 1, n)) for i in range(1, n + 1, k)]
    else:
        if not (opts.get("ranges") or "").strip():
            raise ToolError("Enter at least one range, e.g. 1-3, 4-")
        ranges = parse_ranges(opts.get("ranges", ""), n)
    outs = []
    for a, b in ranges:
        with pikepdf.new() as part:
            part.pages.extend(src.pages[a - 1 : b])
            p = out_path(work, files[0], ".pdf", f"p{a}" if a == b else f"p{a}-{b}")
            part.save(p)
            outs.append(p)
    return single_or_zip(outs, work, files[0].stem + "_split.zip")


def _select_pages(files, work, spec: str, keep: bool, tag: str):
    pdf = pike(files[0])
    n = len(pdf.pages)
    chosen = page_set(spec, n)
    if keep:
        new = pikepdf.new()
        for i in chosen:
            new.pages.append(pdf.pages[i])
    else:
        drop = set(chosen)
        if len(drop) >= n:
            raise ToolError("That would remove every page.")
        for i in sorted(drop, reverse=True):
            del pdf.pages[i]
        new = pdf
    out = out_path(work, files[0], ".pdf", tag)
    new.save(out)
    return Result("file", out)


@tool(
    id="remove-pages", name="Remove pages", category="organize", glyph="remove",
    description="Delete the pages you don't need.",
    accept=PDF, options=[{**PAGES_OPT, "label": "Pages to remove", "help": "", "required": True}],
)
def remove_pages(files, opts, work):
    if not (opts.get("pages") or "").strip():
        raise ToolError("Tell me which pages to remove.")
    return _select_pages(files, work, opts["pages"], keep=False, tag="trimmed")


@tool(
    id="extract-pages", name="Extract pages", category="organize", glyph="extract",
    description="Pull selected pages out into a new PDF.",
    accept=PDF, options=[{**PAGES_OPT, "label": "Pages to keep", "help": "", "required": True}],
)
def extract_pages(files, opts, work):
    if not (opts.get("pages") or "").strip():
        raise ToolError("Tell me which pages to extract.")
    return _select_pages(files, work, opts["pages"], keep=True, tag="extract")


@tool(
    id="organize", name="Reorder pages", category="organize", glyph="reorder",
    description="Set a new page order, or flip the whole document.",
    accept=PDF,
    options=[
        {"name": "mode", "label": "Mode", "type": "segmented", "default": "custom",
         "choices": [["custom", "Custom order"], ["reverse", "Reverse"], ["odd-even", "Odd, then even"]]},
        {"name": "order", "label": "New order", "type": "text", "placeholder": "3, 1, 2, 4-",
         "help": "Pages not listed are appended in their original order.", "default": "", "when": {"mode": "custom"}},
    ],
)
def organize(files, opts, work):
    pdf = pike(files[0])
    n = len(pdf.pages)
    mode = opts.get("mode", "custom")
    if mode == "reverse":
        order = list(range(n - 1, -1, -1))
    elif mode == "odd-even":
        order = list(range(0, n, 2)) + list(range(1, n, 2))
    else:
        order = page_set(opts.get("order", ""), n) if (opts.get("order") or "").strip() else []
        order += [i for i in range(n) if i not in set(order)]
    new = pikepdf.new()
    for i in order:
        new.pages.append(pdf.pages[i])
    out = out_path(work, files[0], ".pdf", "reordered")
    new.save(out)
    return Result("file", out)


@tool(
    id="rotate", name="Rotate PDF", category="organize", glyph="rotate",
    description="Turn pages 90°, 180° or 270° — all of them or just a few.",
    accept=PDF, multiple=True,
    options=[
        {"name": "angle", "label": "Angle", "type": "segmented", "default": "90",
         "choices": [["90", "90° ↻"], ["180", "180°"], ["270", "90° ↺"]]},
        PAGES_OPT,
    ],
)
def rotate(files, opts, work):
    angle = int(opts.get("angle") or 90)

    def one(f):
        pdf = pike(f)
        for i in page_set(opts.get("pages", ""), len(pdf.pages)):
            pdf.pages[i].rotate(angle, relative=True)
        out = out_path(work, f, ".pdf", "rotated")
        pdf.save(out)
        return out

    return for_each(files, work, one, "rotated.zip")


# --------------------------------------------------------------------------- #
# Optimize
# --------------------------------------------------------------------------- #




@tool(
    id="compress", name="Compress PDF", category="optimize", glyph="compress", popular=True, keywords="reduce shrink smaller",
    description="Reduce file size while keeping the quality you need.",
    accept=PDF, multiple=True,
    options=[
        {"name": "level", "label": "Compression", "type": "cards", "default": "ebook", "choices": [
            ["screen", "Extreme", "Smallest file, 100 dpi images"],
            ["ebook", "Recommended", "Good quality, 150 dpi"],
            ["printer", "Light", "High quality, 225 dpi"],
        ]},
    ],
)
def compress(files, opts, work):
    level = opts.get("level", "ebook")
    if level not in {"screen", "ebook", "printer"}:
        level = "ebook"
    total_in = sum(f.stat().st_size for f in files)

    def one(f):
        out = out_path(work, f, ".pdf", "compressed")
        _gs_downsample(f, out, {"screen": 100, "ebook": 150, "printer": 225}[level],
                       {"screen": "/screen", "ebook": "/ebook", "printer": "/printer"}[level])
        if out.stat().st_size >= f.stat().st_size:  # never hand back a bigger file
            shutil.copy(f, out)
        return out

    res = for_each(files, work, one, "compressed.zip")
    res.original_size = total_in
    return res


@tool(
    id="repair", name="Repair PDF", category="optimize", glyph="repair",
    description="Rebuild damaged or malformed PDFs so they open again.",
    accept=PDF, multiple=True,
)
def repair(files, opts, work):
    def one(f):
        out = out_path(work, f, ".pdf", "repaired")
        try:
            with pikepdf.open(f, attempt_recovery=True) as pdf:
                pdf.save(out)
        except Exception:  # noqa: BLE001 - next: MuPDF's repairing parser, then a full Ghostscript rewrite
            try:
                fitz.open(f).save(out, garbage=3, deflate=True)
            except Exception:  # noqa: BLE001
                run(["gs", "-sDEVICE=pdfwrite", "-dNOPAUSE", "-dBATCH", "-dQUIET", "-dSAFER", f"-sOutputFile={out}", str(f)])
        try:
            rec = fitz.open(out)
            ok = rec.page_count > 0 and any(p.get_text().strip() or p.get_images() or p.get_drawings() for p in rec)
        except Exception:  # noqa: BLE001
            ok = False
        if not ok:
            raise ToolError(f"{f.name} is too damaged to recover any pages.")
        return out

    return for_each(files, work, one, "repaired.zip")


OCR_LANGS = [["eng", "English"], ["deu", "German"], ["fra", "French"], ["spa", "Spanish"],
             ["ita", "Italian"], ["por", "Portuguese"], ["hin", "Hindi"]]


@tool(
    id="ocr", name="OCR PDF", category="optimize", glyph="ocr",
    description="Make scanned PDFs searchable and selectable with OCRmyPDF.",
    accept=PDF + [".jpg", ".jpeg", ".png", ".tif", ".tiff"],
    options=[
        {"name": "lang", "label": "Document language", "type": "select", "default": "eng", "choices": OCR_LANGS},
        {"name": "deskew", "label": "Straighten crooked scans", "type": "checkbox", "default": True},
        {"name": "force", "label": "Redo OCR on pages that already have text", "type": "checkbox", "default": False},
    ],
)
def ocr(files, opts, work):
    f = files[0]
    if f.suffix.lower() != ".pdf":
        conv = work / "img.pdf"
        img = fitz.open(f)
        conv.write_bytes(img.convert_to_pdf())
        f_in = conv
    else:
        f_in = f
    lang = opts.get("lang", "eng")
    if lang not in {c[0] for c in OCR_LANGS}:
        lang = "eng"
    out = out_path(work, f, ".pdf", "ocr")
    cmd = ["ocrmypdf", "-l", lang, "--output-type", "pdf", "--optimize", "1"]
    cmd.append("--force-ocr" if opts.get("force") else "--skip-text")
    if opts.get("deskew") and opts.get("force"):
        cmd.append("--deskew")
    cmd += [str(f_in), str(out)]
    run(cmd, timeout=1800)
    return Result("file", out)


@tool(
    id="grayscale", name="Grayscale PDF", category="optimize", glyph="gray",
    description="Strip colour for cheaper printing and smaller files.",
    accept=PDF, multiple=True,
)
def grayscale(files, opts, work):
    def one(f):
        out = out_path(work, f, ".pdf", "gray")
        ghostscript(f, out, "-sColorConversionStrategy=Gray", "-dProcessColorModel=/DeviceGray")
        return out

    return for_each(files, work, one, "grayscale.zip")


# --------------------------------------------------------------------------- #
# Convert to PDF
# --------------------------------------------------------------------------- #




def _office_to_pdf(files, opts, work):
    return single_or_zip(soffice(files, work), work, "converted.zip")


for _id, _name, _accept, _desc in [
    ("word-to-pdf", "Word to PDF", OFFICE_WORD, "DOC, DOCX, ODT and RTF to pixel-faithful PDF."),
    ("powerpoint-to-pdf", "PowerPoint to PDF", OFFICE_PPT, "Turn slide decks into shareable PDFs."),
    ("excel-to-pdf", "Excel to PDF", OFFICE_XLS, "Spreadsheets to print-ready PDF."),
]:
    tool(id=_id, name=_name, category="to-pdf", description=_desc, accept=_accept,
         multiple=True, glyph=_id.split("-")[0])(_office_to_pdf)


PAGE_SIZES = {"a4": fitz.paper_rect("a4"), "letter": fitz.paper_rect("letter")}


@tool(
    id="image-to-pdf", name="Images to PDF", category="to-pdf", glyph="image",
    description="JPG, PNG, WebP, TIFF and more — one image per page.",
    accept=IMAGES, multiple=True, reorder=True,
    options=[
        {"name": "size", "label": "Page size", "type": "segmented", "default": "fit",
         "choices": [["fit", "Fit image"], ["a4", "A4"], ["letter", "Letter"]]},
        {"name": "orientation", "label": "Orientation", "type": "segmented", "default": "auto",
         "choices": [["auto", "Auto"], ["portrait", "Portrait"], ["landscape", "Landscape"]],
         "when": {"size": ["a4", "letter"]}},
        {"name": "margin", "label": "Margin (mm)", "type": "number", "min": 0, "max": 50, "default": 10,
         "when": {"size": ["a4", "letter"]}},
    ],
)
def image_to_pdf(files, opts, work):
    out_doc = fitz.open()
    size = opts.get("size", "fit")
    margin = float(opts.get("margin") or 0) * 72 / 25.4
    for f in files:
        try:
            img = fitz.open(f)
            pdf_bytes = img.convert_to_pdf()
        except Exception as e:  # noqa: BLE001
            raise ToolError(f"Can't read image {f.name}") from e
        img_pdf = fitz.open("pdf", pdf_bytes)
        if size == "fit":
            out_doc.insert_pdf(img_pdf)
            continue
        ir = img_pdf[0].rect
        pr = PAGE_SIZES[size]
        orient = opts.get("orientation", "auto")
        landscape = orient == "landscape" or (orient == "auto" and ir.width > ir.height)
        w, h = (pr.height, pr.width) if landscape else (pr.width, pr.height)
        page = out_doc.new_page(width=w, height=h)
        page.show_pdf_page(fitz.Rect(margin, margin, w - margin, h - margin), img_pdf, 0, keep_proportion=True)
    out = work / "out" / (files[0].stem + ".pdf" if len(files) == 1 else "images.pdf")
    out_doc.save(out, deflate=True, garbage=3)
    return Result("file", out)


DOC_CSS = """
body { font-family: serif; font-size: 11pt; line-height: 1.45; color: #1c1a17; }
h1 { font-size: 22pt; margin: 0 0 8pt; } h2 { font-size: 16pt; margin: 14pt 0 6pt; }
h3 { font-size: 13pt; margin: 12pt 0 4pt; }
p, li { margin: 0 0 6pt; }
code, pre { font-family: monospace; font-size: 9.5pt; background-color: #f2eee6; }
pre { padding: 6pt; }
blockquote { margin-left: 12pt; color: #555; font-style: italic; }
table { border-collapse: collapse; } td, th { border: 0.5pt solid #999; padding: 3pt 5pt; }
th { background-color: #eee; }
a { color: #b8321a; }
"""


def html_to_pdf_story(html_text: str, out: Path, base: Path | None = None) -> None:
    archive = fitz.Archive(str(base)) if base else None
    story = fitz.Story(html=html_text, user_css=DOC_CSS, archive=archive)
    writer = fitz.DocumentWriter(str(out))
    mediabox = fitz.paper_rect("a4")
    where = mediabox + (54, 60, -54, -60)
    more = True
    while more:
        dev = writer.begin_page(mediabox)
        more, _ = story.place(where)
        story.draw(dev)
        writer.end_page()
    writer.close()


@tool(
    id="markdown-to-pdf", name="Markdown to PDF", category="to-pdf", glyph="md",
    description="Typeset .md files into tidy A4 PDF documents.",
    accept=[".md", ".markdown", ".txt"], multiple=True,
)
def markdown_to_pdf(files, opts, work):
    import markdown as mdlib

    def one(f):
        text = f.read_text(encoding="utf-8", errors="replace")
        body = mdlib.markdown(text, extensions=["tables", "fenced_code", "sane_lists"])
        out = out_path(work, f, ".pdf")
        html_to_pdf_story(f"<body>{body}</body>", out, f.parent)
        return out

    return for_each(files, work, one, "markdown_pdfs.zip")


@tool(
    id="html-to-pdf", name="HTML to PDF", category="to-pdf", glyph="globe",
    description="Render .html files to PDF with Chromium, CSS and all.",
    accept=[".html", ".htm"], multiple=True,
)
def html_to_pdf(files, opts, work):
    def one(f):
        tmp = work / f"{f.stem}.chrome.pdf"
        chromium(["--no-pdf-header-footer", "--allow-file-access-from-files", f"--print-to-pdf={tmp}", f.as_uri()], work)
        if not tmp.exists():
            return soffice([f], work, "pdf:writer_web_pdf_Export")[0]
        dest = out_path(work, f, ".pdf")
        shutil.move(tmp, dest)
        return dest

    return for_each(files, work, one, "html_pdfs.zip")


@tool(
    id="pdf-to-pdfa", name="PDF to PDF/A", category="to-pdf", glyph="archive",
    description="Convert to PDF/A-2b for long-term archiving.",
    accept=PDF,
)
def pdf_to_pdfa(files, opts, work):
    out = out_path(work, files[0], ".pdf", "pdfa")
    run(["ocrmypdf", "--skip-text", "--tesseract-timeout", "0", "--output-type", "pdfa-2",
         str(files[0]), str(out)], timeout=900)
    return Result("file", out)


# --------------------------------------------------------------------------- #
# Convert from PDF
# --------------------------------------------------------------------------- #


@tool(
    id="pdf-to-image", name="PDF to JPG / PNG", category="from-pdf", glyph="image",
    description="Every page as a crisp image, bundled into a ZIP.",
    accept=PDF,
    options=[
        {"name": "format", "label": "Format", "type": "segmented", "default": "jpg",
         "choices": [["jpg", "JPG"], ["png", "PNG"], ["webp", "WebP"]]},
        {"name": "dpi", "label": "Resolution", "type": "segmented", "default": "150",
         "choices": [["72", "72 dpi"], ["150", "150 dpi"], ["300", "300 dpi"]]},
        PAGES_OPT,
    ],
)
def pdf_to_image(files, opts, work):
    doc = open_pdf(files[0])
    fmt = opts.get("format", "jpg") if opts.get("format") in {"jpg", "png", "webp"} else "jpg"
    dpi = min(600, max(36, int(opts.get("dpi") or 150)))
    outs = []
    for i in page_set(opts.get("pages", ""), doc.page_count):
        pix = doc[i].get_pixmap(dpi=dpi, alpha=False)
        p = out_path(work, files[0], f".{fmt}", f"page{i + 1:03d}")
        if fmt == "webp":
            from PIL import Image

            Image.frombytes("RGB", (pix.width, pix.height), pix.samples).save(p, "WEBP", quality=88)
        elif fmt == "jpg":
            pix.save(p, jpg_quality=90)
        else:
            pix.save(p)
        outs.append(p)
    return single_or_zip(outs, work, files[0].stem + "_images.zip")


@tool(
    id="pdf-to-word", name="PDF to Word", category="from-pdf", glyph="word", popular=True, keywords="docx editable",
    description="Editable DOCX — flowing text with tables, or an exact-layout copy with shapes and charts.",
    accept=PDF,
    options=[{"name": "mode", "label": "Conversion", "type": "cards", "default": "flow", "choices": [
        ["flow", "Editable text", "Paragraphs and tables you can retype and reflow"],
        ["exact", "Exact layout", "Looks identical — text boxes, shapes and charts kept in place"]]}],
)
def pdf_to_word(files, opts, work):
    open_pdf(files[0]).close()
    out = out_path(work, files[0], ".docx")
    if opts.get("mode") == "exact":
        produced = soffice([files[0]], work, "docx:MS Word 2007 XML", infilter="writer_pdf_import")[0]
        produced.replace(out)
    else:
        from pdf2docx import Converter

        cv = Converter(str(files[0]))
        try:
            cv.convert(str(out))
        finally:
            cv.close()
    shrink_office_media(out)
    return Result("file", out)


@tool(
    id="pdf-to-powerpoint", name="PDF to PowerPoint", category="from-pdf", glyph="powerpoint",
    description="One slide per page, sized to match your document.",
    accept=PDF,
)
def pdf_to_powerpoint(files, opts, work):
    from pptx import Presentation
    from pptx.util import Emu

    doc = open_pdf(files[0])
    prs = Presentation()
    r = doc[0].rect
    prs.slide_width = Emu(int(r.width * 12700))
    prs.slide_height = Emu(int(r.height * 12700))
    blank = prs.slide_layouts[6]
    for page in doc:
        pix = page.get_pixmap(dpi=150, alpha=False)
        slide = prs.slides.add_slide(blank)
        slide.shapes.add_picture(io.BytesIO(pix.tobytes("jpg", jpg_quality=88)), 0, 0, prs.slide_width, prs.slide_height)
        notes = page.get_text().strip()
        if notes:
            slide.notes_slide.notes_text_frame.text = notes[:5000]
    out = out_path(work, files[0], ".pptx")
    prs.save(out)
    return Result("file", out)


@tool(
    id="pdf-to-excel", name="PDF to Excel", category="from-pdf", glyph="excel",
    description="Detect tables and drop each one into its own sheet.",
    accept=PDF,
)
def pdf_to_excel(files, opts, work):
    from openpyxl import Workbook
    from openpyxl.styles import Font

    doc = open_pdf(files[0])
    wb = Workbook()
    wb.remove(wb.active)
    for pno, page in enumerate(doc, 1):
        for tno, table in enumerate(page.find_tables().tables, 1):
            ws = wb.create_sheet(f"P{pno} T{tno}"[:31])
            for row in table.extract():
                ws.append([(c or "").strip() if isinstance(c, str) else c for c in row])
            for cell in ws[1]:
                cell.font = Font(bold=True)
    if not wb.sheetnames:  # no tables: fall back to one line per row of text
        ws = wb.create_sheet("Text")
        for pno, page in enumerate(doc, 1):
            for line in page.get_text().splitlines():
                if line.strip():
                    ws.append([pno, line])
    out = out_path(work, files[0], ".xlsx")
    wb.save(out)
    return Result("file", out)


@tool(
    id="pdf-to-text", name="PDF to Text", category="from-pdf", glyph="text",
    description="Plain UTF-8 text, page by page.",
    accept=PDF, multiple=True,
)
def pdf_to_text(files, opts, work):
    def one(f):
        doc = open_pdf(f)
        out = out_path(work, f, ".txt")
        out.write_text("\n\f\n".join(p.get_text() for p in doc), encoding="utf-8")
        return out

    return for_each(files, work, one, "text.zip")


@tool(
    id="extract-images", name="Extract images", category="from-pdf", glyph="extract",
    description="Save every embedded picture at its original resolution.",
    accept=PDF,
)
def extract_images(files, opts, work):
    doc = open_pdf(files[0])
    outs, seen = [], set()
    for pno, page in enumerate(doc, 1):
        for img in page.get_images(full=True):
            xref = img[0]
            if xref in seen:
                continue
            seen.add(xref)
            info = doc.extract_image(xref)
            if not info or info.get("width", 0) < 16:
                continue
            p = work / "out" / f"page{pno:03d}_img{xref}.{info['ext']}"
            p.write_bytes(info["image"])
            outs.append(p)
    if not outs:
        raise ToolError("No embedded images found in this PDF.")
    return Result("file", zip_files(outs, work / "out" / f"{files[0].stem}_images.zip"))


# --------------------------------------------------------------------------- #
# Edit
# --------------------------------------------------------------------------- #


@tool(
    id="watermark", name="Add watermark", category="edit", glyph="watermark",
    description="Stamp text diagonally across every page.",
    accept=PDF, multiple=True,
    options=[
        {"name": "kind", "label": "Watermark", "type": "segmented", "default": "text",
         "choices": [["text", "Text"], ["image", "Image / logo"]]},
        {"name": "text", "label": "Watermark text", "type": "text", "default": "CONFIDENTIAL", "when": {"kind": "text"}},
        {"name": "asset", "label": "Logo image", "type": "file", "accept": "image/*", "when": {"kind": "image"}},
        {"name": "scale", "label": "Logo width", "type": "range", "min": 5, "max": 100, "default": 40, "unit": "%",
         "when": {"kind": "image"}},
        {"name": "size", "label": "Font size", "type": "range", "min": 12, "max": 140, "default": 64, "when": {"kind": "text"}},
        {"name": "opacity", "label": "Opacity", "type": "range", "min": 5, "max": 100, "default": 35, "unit": "%"},
        {"name": "angle", "label": "Angle", "type": "range", "min": -90, "max": 90, "default": 45, "unit": "°", "when": {"kind": "text"}},
        {"name": "color", "label": "Colour", "type": "color", "default": "#c8361d", "when": {"kind": "text"}},
        PAGES_OPT,
    ],
)
def watermark(files, opts, work):
    if opts.get("kind") == "image":
        return _image_watermark(files, opts, work)
    text = (opts.get("text") or "").strip()
    if not text:
        raise ToolError("Watermark text can't be empty.")
    size = float(opts.get("size") or 64)
    opacity = float(opts.get("opacity") or 35) / 100
    angle = float(opts.get("angle") or 0)
    color = hex_color(opts.get("color"))

    def one(f):
        doc = open_pdf(f)
        tw = fitz.get_text_length(text, fontname="hebo", fontsize=size)
        for i in page_set(opts.get("pages", ""), doc.page_count):
            page = doc[i]
            r = page.rect
            c = fitz.Point((r.x0 + r.x1) / 2, (r.y0 + r.y1) / 2) * page.derotation_matrix
            origin = fitz.Point(c.x - tw / 2, c.y + size * 0.35)
            page.insert_text(origin, text, fontsize=size, fontname="hebo", color=color,
                             fill_opacity=opacity, stroke_opacity=opacity,
                             morph=(c, fitz.Matrix(angle + page.rotation)), overlay=True)
        out = out_path(work, f, ".pdf", "watermarked")
        doc.save(out, garbage=3, deflate=True)
        return out

    return for_each(files, work, one, "watermarked.zip")


def _image_watermark(files, opts, work):
    from PIL import Image

    asset = opts.get("asset")
    if not asset:
        raise ToolError("Choose a logo image.")
    im = Image.open(asset).convert("RGBA")
    alpha = im.getchannel("A").point(lambda a: int(a * num(opts, "opacity", 18) / 100))
    im.putalpha(alpha)
    buf = io.BytesIO()
    im.save(buf, "PNG")
    data = buf.getvalue()
    frac = num(opts, "scale", 40) / 100

    def one(f):
        doc = open_pdf(f)
        for i in page_set(opts.get("pages", ""), doc.page_count):
            page = doc[i]
            r = page.rect
            w = r.width * frac
            h = w * im.height / im.width
            box = fitz.Rect((r.width - w) / 2, (r.height - h) / 2, (r.width + w) / 2, (r.height + h) / 2)
            page.insert_image(box * page.derotation_matrix, stream=data, rotate=-page.rotation % 360, overlay=True)
        out = out_path(work, f, ".pdf", "watermarked")
        doc.save(out, garbage=3, deflate=True)
        return out

    return for_each(files, work, one, "watermarked.zip")


POSITIONS = [["bottom-center", "Bottom centre"], ["bottom-right", "Bottom right"], ["bottom-left", "Bottom left"],
             ["top-center", "Top centre"], ["top-right", "Top right"], ["top-left", "Top left"]]


@tool(
    id="page-numbers", name="Add page numbers", category="edit", glyph="numbers",
    description="Number pages with your choice of position and format.",
    accept=PDF, multiple=True,
    options=[
        {"name": "position", "label": "Position", "type": "select", "default": "bottom-center", "choices": POSITIONS},
        {"name": "format", "label": "Format", "type": "segmented", "default": "n",
         "choices": [["n", "1"], ["page-n", "Page 1"], ["n-of", "1 / 9"], ["page-n-of", "Page 1 of 9"]]},
        {"name": "start", "label": "Start at", "type": "number", "min": 0, "default": 1},
        {"name": "size", "label": "Font size", "type": "range", "min": 6, "max": 24, "default": 10},
        PAGES_OPT,
    ],
)
def page_numbers(files, opts, work):
    fmt = opts.get("format", "n")
    start = int(opts.get("start") or 1)
    size = float(opts.get("size") or 10)
    pos = opts.get("position", "bottom-center")
    vert, horiz = pos.split("-")

    def label(k, total):
        return {"n": f"{k}", "page-n": f"Page {k}", "n-of": f"{k} / {total}",
                "page-n-of": f"Page {k} of {total}"}.get(fmt, f"{k}")

    def one(f):
        doc = open_pdf(f)
        idx = page_set(opts.get("pages", ""), doc.page_count)
        total = start + len(idx) - 1
        for k, i in enumerate(idx, start):
            page = doc[i]
            r = page.rect  # visible (rotated) coordinates
            s = label(k, total)
            w = fitz.get_text_length(s, fontname="helv", fontsize=size)
            m = 28
            x = {"left": m, "center": (r.width - w) / 2, "right": r.width - m - w}[horiz]
            y = m + size if vert == "top" else r.height - m
            p = fitz.Point(x, y) * page.derotation_matrix
            page.insert_text(p, s, fontsize=size, fontname="helv", color=(0.15, 0.15, 0.15),
                             rotate=page.rotation)
        out = out_path(work, f, ".pdf", "numbered")
        doc.save(out, garbage=3, deflate=True)
        return out

    return for_each(files, work, one, "numbered.zip")


@tool(
    id="crop", name="Crop PDF", category="edit", glyph="crop",
    description="Trim margins from every page by a fixed amount.",
    accept=PDF,
    options=[
        {"name": "top", "label": "Top (mm)", "type": "number", "min": 0, "default": 10},
        {"name": "right", "label": "Right (mm)", "type": "number", "min": 0, "default": 10},
        {"name": "bottom", "label": "Bottom (mm)", "type": "number", "min": 0, "default": 10},
        {"name": "left", "label": "Left (mm)", "type": "number", "min": 0, "default": 10},
        PAGES_OPT,
    ],
)
def crop(files, opts, work):
    mm = 72 / 25.4
    t, r, b, l = (float(opts.get(k) or 0) * mm for k in ("top", "right", "bottom", "left"))
    doc = open_pdf(files[0])
    for i in page_set(opts.get("pages", ""), doc.page_count):
        page = doc[i]
        mb = page.mediabox
        # cropbox is in unrotated PDF space where y grows upward; fitz handles the flip.
        rect = fitz.Rect(mb.x0 + l, mb.y0 + t, mb.x1 - r, mb.y1 - b)
        if rect.is_empty or rect.width < 20 or rect.height < 20:
            raise ToolError(f"Margins too large for page {i + 1}.")
        page.set_cropbox(rect)
    out = out_path(work, files[0], ".pdf", "cropped")
    doc.save(out, garbage=3, deflate=True)
    return Result("file", out)


@tool(
    id="flatten", name="Flatten PDF", category="edit", glyph="flatten",
    description="Burn form fields and annotations into the page so they can't change.",
    accept=PDF, multiple=True,
)
def flatten(files, opts, work):
    def one(f):
        doc = open_pdf(f)
        doc.bake(annots=True, widgets=True)
        out = out_path(work, f, ".pdf", "flat")
        doc.save(out, garbage=3, deflate=True)
        return out

    return for_each(files, work, one, "flattened.zip")


@tool(
    id="compare", name="Compare PDF", category="edit", glyph="compare",
    description="Side-by-side text diff of two versions of a document.",
    accept=PDF, multiple=True, min_files=2, reorder=True,
)
def compare(files, opts, work):
    if len(files) != 2:
        raise ToolError("Upload exactly two PDFs: the original first, the revision second.")
    a, b = ([ln for p in open_pdf(f) for ln in p.get_text().splitlines() if ln.strip()] for f in files)
    table = difflib.HtmlDiff(wrapcolumn=70).make_table(a, b, html.escape(files[0].name), html.escape(files[1].name),
                                                       context=True, numlines=2)
    sm = difflib.SequenceMatcher(None, a, b, autojunk=False)
    ratio = round(sm.ratio() * 100, 1)
    adds = sum(j2 - j1 for op, i1, i2, j1, j2 in sm.get_opcodes() if op in ("insert", "replace"))
    dels = sum(i2 - i1 for op, i1, i2, j1, j2 in sm.get_opcodes() if op in ("delete", "replace"))
    doc = f"""<!doctype html><meta charset="utf-8"><title>Comparison</title>
<style>
body{{font-family:Figtree,system-ui,sans-serif;margin:0;padding:24px;color:#1c1a17;background:#fbf8f2}}
h1{{font-family:Fraunces,Georgia,serif;font-weight:500;margin:0 0 4px}}
.stats{{margin:0 0 20px;color:#6b645a}} .stats b{{color:#1c1a17}}
table.diff{{border-collapse:collapse;width:100%;font:12px/1.5 'JetBrains Mono',monospace;background:#fff}}
.diff td{{padding:2px 8px;vertical-align:top;white-space:pre-wrap;word-break:break-word}}
.diff th{{background:#efe9dd;padding:6px;text-align:left}} .diff_header{{color:#9a9184;background:#f6f2ea;text-align:right}}
.diff_next{{display:none}} .diff_add{{background:#dcefd9}} .diff_chg{{background:#fbedc5}} .diff_sub{{background:#f7d9d2}}
</style><h1>Document comparison</h1>
<p class="stats"><b>{ratio}%</b> similar · <b>{adds}</b> lines added or changed · <b>{dels}</b> removed</p>
{table if (adds or dels) else '<p>The documents contain identical text.</p>'}"""
    out = work / "out" / "comparison.html"
    out.write_text(doc, encoding="utf-8")
    return Result("html", out, html=doc)


# --------------------------------------------------------------------------- #
# Security
# --------------------------------------------------------------------------- #


@tool(
    id="protect", name="Protect PDF", category="security", glyph="lock",
    description="Encrypt with AES-256 and a password.",
    accept=PDF,
    options=[
        {"name": "password", "label": "Password", "type": "password", "required": True, "default": ""},
        {"name": "no_print", "label": "Block printing", "type": "checkbox", "default": False},
        {"name": "no_copy", "label": "Block copying text", "type": "checkbox", "default": False},
        {"name": "no_modify", "label": "Block editing", "type": "checkbox", "default": False},
    ],
)
def protect(files, opts, work):
    pw = opts.get("password") or ""
    if len(pw) < 4:
        raise ToolError("Use a password of at least 4 characters.")
    allow = pikepdf.Permissions(
        print_lowres=not opts.get("no_print"), print_highres=not opts.get("no_print"),
        extract=not opts.get("no_copy"), modify_other=not opts.get("no_modify"),
        modify_annotation=not opts.get("no_modify"), modify_form=not opts.get("no_modify"),
        modify_assembly=not opts.get("no_modify"),
    )
    out = out_path(work, files[0], ".pdf", "protected")
    try:
        with pikepdf.open(files[0]) as pdf:
            pdf.save(out, encryption=pikepdf.Encryption(user=pw, owner=pw, R=6, allow=allow))
    except pikepdf.PasswordError as e:
        raise ToolError("This PDF is already password protected.") from e
    return Result("file", out)


@tool(
    id="unlock", name="Unlock PDF", category="security", glyph="unlock",
    description="Remove the password and restrictions from a PDF you own.",
    accept=PDF,
    options=[{"name": "password", "label": "Current password", "type": "password", "default": "",
              "help": "Leave empty if it only has editing restrictions."}],
)
def unlock(files, opts, work):
    out = out_path(work, files[0], ".pdf", "unlocked")
    try:
        with pikepdf.open(files[0], password=opts.get("password") or "") as pdf:
            pdf.save(out)
    except pikepdf.PasswordError as e:
        raise ToolError("Wrong password.") from e
    return Result("file", out)


@tool(
    id="redact", name="Redact PDF", category="security", glyph="redact",
    description="Permanently black out words, names or numbers — the text is removed, not hidden.",
    accept=PDF,
    options=[
        {"name": "terms", "label": "Words or phrases", "type": "textarea", "required": True, "default": "",
         "placeholder": "One per line", "help": "Matching is case-insensitive."},
        {"name": "regex", "label": "Also redact emails & phone numbers", "type": "checkbox", "default": False},
    ],
)
def redact(files, opts, work):
    terms = [t.strip() for t in (opts.get("terms") or "").splitlines() if t.strip()]
    patterns = []
    if opts.get("regex"):
        patterns = [re.compile(r"[\w.+-]+@[\w-]+\.[\w.-]+"), re.compile(r"\+?\d[\d\s().-]{7,}\d")]
    if not terms and not patterns:
        raise ToolError("Add at least one word to redact.")
    doc = open_pdf(files[0])
    hits = 0
    for page in doc:
        found = []
        for t in terms:
            found += page.search_for(t)
        if patterns:
            text = page.get_text()
            for pat in patterns:
                for m in set(pat.findall(text)):
                    found += page.search_for(m)
        for r in found:
            page.add_redact_annot(r, fill=(0, 0, 0))
        hits += len(found)
        if found:
            page.apply_redactions()
    if not hits:
        raise ToolError("None of those terms were found in the document.")
    doc.scrub(metadata=False)
    out = out_path(work, files[0], ".pdf", "redacted")
    doc.save(out, garbage=4, deflate=True)
    return Result("file", out)


@tool(
    id="sanitize", name="Remove metadata", category="security", glyph="sanitize",
    description="Strip author, software, XMP, scripts and hidden data.",
    accept=PDF, multiple=True,
)
def sanitize(files, opts, work):
    def one(f):
        doc = open_pdf(f)
        doc.scrub(metadata=True, xml_metadata=True, javascript=True, embedded_files=True,
                  attached_files=True, thumbnails=True, reset_fields=False, redactions=False)
        doc.set_metadata({})
        out = out_path(work, f, ".pdf", "clean")
        doc.save(out, garbage=4, deflate=True, clean=True)
        return out

    return for_each(files, work, one, "sanitized.zip")


# --------------------------------------------------------------------------- #
# PDF — size & page geometry
# --------------------------------------------------------------------------- #

from .registry import fmt_size, num, target_bytes  # noqa: E402

TARGET_OPTS = [
    {"name": "target", "label": "Target size", "type": "number", "min": 1, "default": 1, "half": True, "required": True},
    {"name": "unit", "label": "Unit", "type": "segmented", "default": "mb", "half": True,
     "choices": [["kb", "KB"], ["mb", "MB"]]},
]


def _gs_downsample(src: Path, dest: Path, dpi: int, preset: str | None = None) -> None:
    preset = preset or ("/ebook" if dpi >= 110 else "/screen")  # /screen also means harsher JPEG quantization
    ghostscript(src, dest, f"-dPDFSETTINGS={preset}", "-dDetectDuplicateImages=true",
                "-dDownsampleColorImages=true", "-dDownsampleGrayImages=true", "-dDownsampleMonoImages=true",
                f"-dColorImageResolution={dpi}", f"-dGrayImageResolution={dpi}", f"-dMonoImageResolution={dpi * 2}",
                "-dColorImageDownsampleThreshold=1.0", "-dGrayImageDownsampleThreshold=1.0",
                "-dColorImageDownsampleType=/Bicubic", "-dGrayImageDownsampleType=/Bicubic",
                "-dCompressFonts=true", "-dSubsetFonts=true")


def _rasterize(src: Path, dest: Path, dpi: int, quality: int, gray: bool) -> None:
    doc = open_pdf(src)
    out = fitz.open()
    for page in doc:
        pix = page.get_pixmap(dpi=dpi, colorspace=fitz.csGRAY if gray else fitz.csRGB, alpha=False)
        p = out.new_page(width=page.rect.width, height=page.rect.height)
        p.insert_image(p.rect, stream=pix.tobytes("jpg", jpg_quality=quality))
    out.save(dest, garbage=4, deflate=True)


@tool(
    id="compress-pdf-to-size", name="Compress PDF to size", category="optimize", glyph="target", popular=True,
    description="Hit an exact limit — e.g. under 200 KB or 2 MB for upload portals.",
    keywords="kb mb limit target reduce exact",
    accept=PDF,
    options=TARGET_OPTS + [
        {"name": "raster", "label": "If needed, convert pages to images to reach the target", "type": "checkbox", "default": True,
         "help": "Last resort: text stops being selectable."},
        {"name": "gray", "label": "Grayscale when rasterizing", "type": "checkbox", "default": False, "when": {"raster": True}},
        {"name": "strong", "label": "Allow strong quality loss for very small targets", "type": "checkbox", "default": False,
         "help": "Off keeps text legible (≥ 72 dpi); the target may then be missed."},
    ],
)
def compress_pdf_to_size(files, opts, work):
    src = files[0]
    goal = target_bytes(opts)
    if goal <= 0:
        raise ToolError("Enter a target size.")
    orig = src.stat().st_size
    out = out_path(work, src, ".pdf", "sized")
    best: Path | None = None
    if orig <= goal:
        shutil.copy(src, out)
        return Result("file", out, original_size=orig, meta={"target-met": "1", "note": "Already under the target."})
    strong = bool(opts.get("strong"))
    for i, dpi in enumerate([200, 150, 120, 96, 80, 72] + ([60, 50] if strong else [])):
        cand = work / f"gs{i}.pdf"
        try:
            _gs_downsample(src, cand, dpi)
        except ToolError:
            continue
        if best is None or cand.stat().st_size < best.stat().st_size:
            best = cand
        if cand.stat().st_size <= goal:
            break
    if (best is None or best.stat().st_size > goal) and opts.get("raster"):
        gray = bool(opts.get("gray"))
        steps = [(150, 72), (125, 62), (110, 55), (96, 50), (85, 45), (72, 40)] + ([(60, 32), (50, 25)] if strong else [])
        for i, (dpi, q) in enumerate(steps):
            cand = work / f"r{i}.pdf"
            _rasterize(src, cand, dpi, q, gray)
            if best is None or cand.stat().st_size < best.stat().st_size:
                best = cand
            if cand.stat().st_size <= goal:
                break
    if best is None or best.stat().st_size >= orig:
        shutil.copy(src, out)
    else:
        shutil.copy(best, out)
    met = out.stat().st_size <= goal
    note = "" if met else (f"Smallest legible result is {fmt_size(out.stat().st_size)} (target {fmt_size(goal)}). "
                           + ("" if opts.get("strong") else "Tick “Allow strong quality loss” to go smaller."))
    return Result("file", out, original_size=orig, meta={"target-met": "1" if met else "0", "note": note})


PAPER = [["a4", "A4"], ["a3", "A3"], ["a5", "A5"], ["b5", "B5"], ["letter", "Letter"], ["legal", "Legal"],
         ["tabloid", "Tabloid"], ["custom", "Custom (mm)"]]


def _paper(opts: dict, src_rect: fitz.Rect | None = None) -> tuple[float, float]:
    key = opts.get("paper", "a4")
    mm = 72 / 25.4
    if key == "custom":
        w, h = num(opts, "width", 210) * mm, num(opts, "height", 297) * mm
        if w < 20 or h < 20:
            raise ToolError("Custom size is too small.")
        return w, h
    r = fitz.paper_rect(key)
    w, h = r.width, r.height
    orient = opts.get("orientation", "auto")
    want_land = orient == "landscape" or (orient == "auto" and src_rect is not None and src_rect.width > src_rect.height)
    if want_land != (w > h):
        w, h = h, w
    return w, h


PAPER_OPTS = [
    {"name": "paper", "label": "Page size", "type": "select", "default": "a4", "choices": PAPER},
    {"name": "width", "label": "Width (mm)", "type": "number", "min": 20, "default": 210, "half": True, "when": {"paper": "custom"}},
    {"name": "height", "label": "Height (mm)", "type": "number", "min": 20, "default": 297, "half": True, "when": {"paper": "custom"}},
    {"name": "orientation", "label": "Orientation", "type": "segmented", "default": "auto",
     "choices": [["auto", "Auto"], ["portrait", "Portrait"], ["landscape", "Landscape"]]},
]


@tool(
    id="resize-pdf", name="Resize PDF pages", category="optimize", glyph="ratio",
    description="Fit every page onto A4, Letter or any custom size and ratio.",
    keywords="page size ratio scale a4 letter dimensions",
    accept=PDF,
    options=PAPER_OPTS + [
        {"name": "margin", "label": "Margin (mm)", "type": "number", "min": 0, "max": 60, "default": 0},
        {"name": "fit", "label": "Fit", "type": "segmented", "default": "fit",
         "choices": [["fit", "Keep proportions"], ["stretch", "Stretch to fill"]]},
    ],
)
def resize_pdf(files, opts, work):
    src = open_pdf(files[0])
    out = fitz.open()
    m = num(opts, "margin") * 72 / 25.4
    for i, page in enumerate(src):
        w, h = _paper(opts, page.rect)
        p = out.new_page(width=w, height=h)
        p.show_pdf_page(fitz.Rect(m, m, w - m, h - m), src, i, keep_proportion=opts.get("fit") != "stretch")
    dest = out_path(work, files[0], ".pdf", "resized")
    out.save(dest, garbage=3, deflate=True)
    return Result("file", dest)


@tool(
    id="n-up", name="Pages per sheet (N-up)", category="organize", glyph="grid",
    description="Print 2, 4, 6 or 9 pages on each sheet to save paper.",
    keywords="nup imposition handout",
    accept=PDF,
    options=[
        {"name": "per", "label": "Pages per sheet", "type": "segmented", "default": "2",
         "choices": [["2", "2"], ["4", "4"], ["6", "6"], ["9", "9"]]},
        {"name": "paper", "label": "Sheet size", "type": "select", "default": "a4", "choices": PAPER[:-1]},
        {"name": "border", "label": "Draw a thin border around each page", "type": "checkbox", "default": False},
    ],
)
def n_up(files, opts, work):
    src = open_pdf(files[0])
    per = int(opts.get("per") or 2)
    cols, rows = {2: (1, 2), 4: (2, 2), 6: (2, 3), 9: (3, 3)}[per]
    r0 = src[0].rect
    landscape_src = r0.width > r0.height
    pr = fitz.paper_rect(opts.get("paper", "a4"))
    W, H = pr.width, pr.height
    if per == 2:  # two portrait pages side by side on a landscape sheet
        cols, rows = (1, 2) if landscape_src else (2, 1)
        W, H = (min(W, H), max(W, H)) if landscape_src else (max(W, H), min(W, H))
    m = 18
    cw, ch = (W - 2 * m) / cols, (H - 2 * m) / rows
    out = fitz.open()
    sheet = None
    for i in range(src.page_count):
        if i % per == 0:
            sheet = out.new_page(width=W, height=H)
        k = i % per
        c, r = k % cols, k // cols
        cell = fitz.Rect(m + c * cw, m + r * ch, m + (c + 1) * cw, m + (r + 1) * ch) + (4, 4, -4, -4)
        sheet.show_pdf_page(cell, src, i)
        if opts.get("border"):
            sheet.draw_rect(cell, color=(0.6, 0.6, 0.6), width=0.5)
    dest = out_path(work, files[0], ".pdf", f"{per}up")
    out.save(dest, garbage=3, deflate=True)
    return Result("file", dest)


@tool(
    id="interleave", name="Alternate & mix", category="organize", glyph="merge",
    description="Combine front-side and back-side scans into one correctly ordered PDF.",
    keywords="duplex odd even scan interleave",
    accept=PDF, multiple=True, min_files=2, reorder=True,
    options=[{"name": "reverse", "label": "Second file is in reverse order (typical for duplex scans)", "type": "checkbox", "default": True}],
)
def interleave(files, opts, work):
    a, b = pike(files[0]), pike(files[1])
    bp = list(b.pages)
    if opts.get("reverse"):
        bp.reverse()
    out = pikepdf.new()
    for i in range(max(len(a.pages), len(bp))):
        if i < len(a.pages):
            out.pages.append(a.pages[i])
        if i < len(bp):
            out.pages.append(bp[i])
    dest = work / "out" / "mixed.pdf"
    out.save(dest)
    return Result("file", dest)


@tool(
    id="insert-blank", name="Insert blank pages", category="organize", glyph="add",
    description="Add empty pages after chosen pages or at the end.",
    accept=PDF,
    options=[
        {"name": "after", "label": "After pages", "type": "text", "placeholder": "e.g. 1, 3 — empty = end", "default": ""},
        {"name": "count", "label": "How many each time", "type": "number", "min": 1, "max": 50, "default": 1},
    ],
)
def insert_blank(files, opts, work):
    doc = open_pdf(files[0])
    n = doc.page_count
    count = max(1, int(num(opts, "count", 1)))
    after = sorted(set(page_set(opts.get("after", ""), n))) if (opts.get("after") or "").strip() else [n - 1]
    for idx in reversed(after):
        r = doc[idx].rect
        for _ in range(count):
            doc.new_page(pno=idx + 1, width=r.width, height=r.height)
    dest = out_path(work, files[0], ".pdf", "blank")
    doc.save(dest, garbage=3, deflate=True)
    return Result("file", dest)


@tool(
    id="remove-blank", name="Remove blank pages", category="organize", glyph="remove",
    description="Detect and drop empty pages from scans automatically.",
    accept=PDF,
    options=[{"name": "sensitivity", "label": "Ink threshold", "type": "range", "min": 1, "max": 50, "default": 10,
              "unit": "‰", "help": "Pages with less ink than this are treated as blank."}],
)
def remove_blank(files, opts, work):
    doc = open_pdf(files[0])
    thr = num(opts, "sensitivity", 10) / 1000
    blank = []
    for i, page in enumerate(doc):
        pix = page.get_pixmap(dpi=40, colorspace=fitz.csGRAY, alpha=False)
        dark = sum(1 for b in pix.samples if b < 200)
        if dark / max(1, len(pix.samples)) < thr:
            blank.append(i)
    if len(blank) == doc.page_count:
        raise ToolError("Every page looks blank — lower the threshold.")
    if not blank:
        raise ToolError("No blank pages found.")
    doc.delete_pages(blank)
    dest = out_path(work, files[0], ".pdf", "noblank")
    doc.save(dest, garbage=3, deflate=True)
    return Result("file", dest, meta={"note": f"Removed {len(blank)} blank page{'s' if len(blank) > 1 else ''}."})


@tool(
    id="split-by-size", name="Split PDF by size", category="organize", glyph="split",
    description="Cut a big PDF into parts that each stay under an email or upload limit.",
    keywords="mb limit email attachment",
    accept=PDF,
    options=[
        {"name": "target", "label": "Max size per part", "type": "number", "min": 0.1, "default": 5, "half": True},
        {"name": "unit", "label": "Unit", "type": "segmented", "default": "mb", "half": True, "choices": [["kb", "KB"], ["mb", "MB"]]},
    ],
)
def split_by_size(files, opts, work):
    limit = target_bytes(opts)
    if limit <= 0:
        raise ToolError("Enter a maximum size.")
    src = pike(files[0])
    parts, cur, outs = [], [], []

    def size_of(pages):
        tmp = pikepdf.new()
        tmp.pages.extend(pages)
        b = io.BytesIO()
        tmp.save(b, compress_streams=True)
        return b.tell()

    for pg in src.pages:
        if cur and size_of(cur + [pg]) > limit:
            parts.append(cur)
            cur = []
        cur.append(pg)
    if cur:
        parts.append(cur)
    for k, part in enumerate(parts, 1):
        p = pikepdf.new()
        p.pages.extend(part)
        dest = out_path(work, files[0], ".pdf", f"part{k}")
        p.save(dest)
        outs.append(dest)
    return single_or_zip(outs, work, files[0].stem + "_parts.zip")


@tool(
    id="linearize", name="Fast web view", category="optimize", glyph="bolt",
    description="Linearize so the first page shows before the whole file downloads.",
    accept=PDF, multiple=True,
)
def linearize(files, opts, work):
    def one(f):
        dest = out_path(work, f, ".pdf", "web")
        run(["qpdf", "--linearize", str(f), str(dest)])
        return dest

    return for_each(files, work, one, "linearized.zip")


# --------------------------------------------------------------------------- #
# PDF — edit extras
# --------------------------------------------------------------------------- #


@tool(
    id="header-footer", name="Header & footer", category="edit", glyph="numbers",
    description="Running headers and footers with page numbers, dates and file names.",
    keywords="bates page number date stamp",
    accept=PDF, multiple=True,
    options=[
        {"name": "header", "label": "Header text", "type": "text", "default": "", "placeholder": "{file}"},
        {"name": "footer", "label": "Footer text", "type": "text", "default": "Page {page} of {total}",
         "help": "Placeholders: {page} {total} {date} {file}"},
        {"name": "align", "label": "Alignment", "type": "segmented", "default": "center",
         "choices": [["left", "Left"], ["center", "Centre"], ["right", "Right"]]},
        {"name": "size", "label": "Font size", "type": "range", "min": 6, "max": 20, "default": 9},
        {"name": "font", "label": "Typeface", "type": "segmented", "default": "helv",
         "choices": [["helv", "Sans"], ["tiro", "Serif"], ["cour", "Mono"]]},
        {"name": "start", "label": "First number", "type": "number", "default": 1, "half": True},
        {"name": "margin", "label": "Margin (mm)", "type": "number", "default": 10, "half": True},
        PAGES_OPT,
    ],
)
def header_footer(files, opts, work):
    import datetime as _dt

    size = num(opts, "size", 9)
    font = opts.get("font") if opts.get("font") in ("helv", "tiro", "cour") else "helv"
    m = num(opts, "margin", 10) * 72 / 25.4
    start = int(num(opts, "start", 1))
    today = _dt.date.today().isoformat()
    if not (opts.get("header") or opts.get("footer")):
        raise ToolError("Enter a header or a footer.")

    def one(f):
        doc = open_pdf(f)
        idx = page_set(opts.get("pages", ""), doc.page_count)
        total = start + len(idx) - 1
        for k, i in enumerate(idx, start):
            page = doc[i]
            r = page.rect
            for text, top in ((opts.get("header"), True), (opts.get("footer"), False)):
                if not text:
                    continue
                s = (text.replace("{page}", str(k)).replace("{total}", str(total))
                     .replace("{date}", today).replace("{file}", f.stem))
                w = fitz.get_text_length(s, fontname=font, fontsize=size)
                x = {"left": m, "center": (r.width - w) / 2, "right": r.width - m - w}.get(opts.get("align"), m)
                y = m + size if top else r.height - m
                page.insert_text(fitz.Point(x, y) * page.derotation_matrix, s, fontsize=size, fontname=font,
                                 color=(0.2, 0.2, 0.2), rotate=page.rotation)
        dest = out_path(work, f, ".pdf", "hf")
        doc.save(dest, garbage=3, deflate=True)
        return dest

    return for_each(files, work, one, "header_footer.zip")


@tool(
    id="pdf-metadata", name="Edit metadata", category="edit", glyph="tag",
    description="Set title, author, subject and keywords shown in readers and search.",
    accept=PDF,
    options=[
        {"name": "title", "label": "Title", "type": "text", "default": ""},
        {"name": "author", "label": "Author", "type": "text", "default": ""},
        {"name": "subject", "label": "Subject", "type": "text", "default": ""},
        {"name": "keywords", "label": "Keywords", "type": "text", "default": ""},
        {"name": "clear", "label": "Clear fields left empty", "type": "checkbox", "default": False},
    ],
)
def pdf_metadata(files, opts, work):
    doc = open_pdf(files[0])
    md = {k: v for k, v in (doc.metadata or {}).items()
          if k in ("title", "author", "subject", "keywords", "creator", "creationDate", "modDate") and v}
    for k in ("title", "author", "subject", "keywords"):
        v = (opts.get(k) or "").strip()
        if v:
            md[k] = v
        elif opts.get("clear"):
            md.pop(k, None)
    md["producer"] = "Folio"
    doc.set_metadata(md)
    dest = out_path(work, files[0], ".pdf")
    doc.save(dest, garbage=3, deflate=True)
    return Result("file", dest)


# --------------------------------------------------------------------------- #
# PDF — conversion extras
# --------------------------------------------------------------------------- #


@tool(
    id="pdf-to-svg", name="PDF to SVG", category="from-pdf", glyph="vector",
    description="Each page as a scalable vector graphic.",
    accept=PDF, options=[PAGES_OPT],
)
def pdf_to_svg(files, opts, work):
    doc = open_pdf(files[0])
    outs = []
    for i in page_set(opts.get("pages", ""), doc.page_count):
        p = out_path(work, files[0], ".svg", f"page{i + 1:03d}")
        svg = doc[i].get_svg_image(text_as_path=False)
        svg = re.sub(r"(<svg[^>]*>)", r'\1<rect width="100%" height="100%" fill="#ffffff"/>', svg, count=1)
        p.write_text(svg, encoding="utf-8")
        outs.append(p)
    return single_or_zip(outs, work, files[0].stem + "_svg.zip")


@tool(
    id="pdf-to-html", name="PDF to HTML", category="from-pdf", glyph="globe",
    description="A self-contained web page with the text and images of your PDF.",
    accept=PDF,
)
def pdf_to_html(files, opts, work):
    doc = open_pdf(files[0])
    body = "\n".join(f'<section class="page" id="p{i + 1}">{p.get_text("xhtml")}</section>' for i, p in enumerate(doc))
    page = (f"<!doctype html><meta charset='utf-8'><title>{html.escape(files[0].stem)}</title>"
            "<style>body{max-width:860px;margin:auto;padding:24px;font-family:Georgia,serif;line-height:1.5}"
            ".page{border-bottom:1px solid #ddd;padding:24px 0}img{max-width:100%}</style>" + body)
    dest = out_path(work, files[0], ".html")
    dest.write_text(page, encoding="utf-8")
    return Result("file", dest)


def _order_quad(pts):
    import numpy as np

    pts = pts.reshape(4, 2).astype("float32")
    s, d = pts.sum(1), np.diff(pts, axis=1).ravel()
    return np.array([pts[s.argmin()], pts[d.argmin()], pts[s.argmax()], pts[d.argmax()]], dtype="float32")


def enhance_scan(path: Path, mode: str, autocrop: bool):
    """Return a PIL image: perspective-corrected (optional) and cleaned up."""
    import cv2
    import numpy as np
    from PIL import Image, ImageOps

    from .tools_image import load_image

    pil = ImageOps.exif_transpose(load_image(path)).convert("RGB")
    img = cv2.cvtColor(np.array(pil), cv2.COLOR_RGB2BGR)
    if autocrop:
        h, w = img.shape[:2]
        k = 800 / max(h, w)
        small = cv2.resize(img, (int(w * k), int(h * k)))
        g = cv2.GaussianBlur(cv2.cvtColor(small, cv2.COLOR_BGR2GRAY), (5, 5), 0)
        edges = cv2.dilate(cv2.Canny(g, 50, 150), np.ones((3, 3), np.uint8))
        cnts, _ = cv2.findContours(edges, cv2.RETR_LIST, cv2.CHAIN_APPROX_SIMPLE)
        for c in sorted(cnts, key=cv2.contourArea, reverse=True)[:5]:
            approx = cv2.approxPolyDP(c, 0.02 * cv2.arcLength(c, True), True)
            if len(approx) == 4 and cv2.contourArea(approx) > 0.25 * small.shape[0] * small.shape[1]:
                q = _order_quad(approx / k)
                wa = int(max(np.linalg.norm(q[2] - q[3]), np.linalg.norm(q[1] - q[0])))
                ha = int(max(np.linalg.norm(q[1] - q[2]), np.linalg.norm(q[0] - q[3])))
                M = cv2.getPerspectiveTransform(q, np.array([[0, 0], [wa, 0], [wa, ha], [0, ha]], dtype="float32"))
                img = cv2.warpPerspective(img, M, (wa, ha))
                break
    if mode == "bw":
        g = cv2.cvtColor(img, cv2.COLOR_BGR2GRAY)
        return Image.fromarray(cv2.adaptiveThreshold(g, 255, cv2.ADAPTIVE_THRESH_GAUSSIAN_C, cv2.THRESH_BINARY, 31, 15))
    if mode == "gray":
        g = cv2.cvtColor(img, cv2.COLOR_BGR2GRAY)
        return Image.fromarray(cv2.createCLAHE(clipLimit=2.0, tileGridSize=(8, 8)).apply(g))
    if mode == "auto":
        lab = cv2.cvtColor(img, cv2.COLOR_BGR2LAB)
        l, a, b = cv2.split(lab)
        l = cv2.createCLAHE(clipLimit=2.0, tileGridSize=(8, 8)).apply(l)
        img = cv2.cvtColor(cv2.merge((l, a, b)), cv2.COLOR_LAB2BGR)
    return Image.fromarray(cv2.cvtColor(img, cv2.COLOR_BGR2RGB))


@tool(
    id="scan-to-pdf", name="Scan to PDF", category="to-pdf", glyph="camera", capture=True, popular=True,
    description="Snap pages with your phone camera — edges straightened, contrast cleaned up.",
    keywords="camera phone document scanner photo",
    accept=IMAGES, multiple=True, reorder=True,
    options=[
        {"name": "mode", "label": "Look", "type": "segmented", "default": "auto",
         "choices": [["auto", "Auto"], ["gray", "Gray"], ["bw", "B&W"], ["color", "Original"]]},
        {"name": "autocrop", "label": "Detect page edges and straighten", "type": "checkbox", "default": True},
        {"name": "paper", "label": "Page size", "type": "segmented", "default": "a4",
         "choices": [["a4", "A4"], ["letter", "Letter"], ["fit", "Fit image"]]},
        {"name": "ocr", "label": "Make text searchable (OCR)", "type": "checkbox", "default": False},
    ],
)
def scan_to_pdf(files, opts, work):
    out = fitz.open()
    for f in files:
        im = enhance_scan(f, opts.get("mode", "auto"), bool(opts.get("autocrop")))
        im.thumbnail((2200, 2200))
        b = io.BytesIO()
        im.save(b, "JPEG", quality=82, optimize=True)
        if opts.get("paper") == "fit":
            page = out.new_page(width=im.width * 72 / 200, height=im.height * 72 / 200)
            page.insert_image(page.rect, stream=b.getvalue())
        else:
            pr = fitz.paper_rect(opts.get("paper", "a4"))
            w, h = (pr.height, pr.width) if im.width > im.height else (pr.width, pr.height)
            page = out.new_page(width=w, height=h)
            page.insert_image(fitz.Rect(18, 18, w - 18, h - 18), stream=b.getvalue(), keep_proportion=True)
    dest = work / "out" / "scan.pdf"
    out.save(dest, garbage=3, deflate=True)
    if opts.get("ocr"):
        ocred = work / "scan_ocr.pdf"
        run(["ocrmypdf", "-l", "eng", "--skip-text", "--optimize", "1", str(dest), str(ocred)], timeout=1800)
        shutil.move(ocred, dest)
    return Result("file", dest)


EBOOKS = [".epub", ".mobi", ".fb2", ".cbz", ".xps", ".oxps"]


@tool(
    id="ebook-to-pdf", name="eBook to PDF", category="to-pdf", glyph="book",
    description="EPUB, MOBI, FB2, comic CBZ and XPS into paginated PDF.",
    accept=EBOOKS, multiple=True,
    options=[{"name": "paper", "label": "Page size", "type": "segmented", "default": "a5",
              "choices": [["a5", "A5"], ["a4", "A4"], ["letter", "Letter"]]},
             {"name": "fontsize", "label": "Text size", "type": "range", "min": 8, "max": 18, "default": 11}],
)
def ebook_to_pdf(files, opts, work):
    def one(f):
        doc = fitz.open(f)
        if doc.is_reflowable:
            r = fitz.paper_rect(opts.get("paper", "a5"))
            doc.layout(width=r.width, height=r.height, fontsize=num(opts, "fontsize", 11))
        dest = out_path(work, f, ".pdf")
        fitz.open("pdf", doc.convert_to_pdf()).save(dest, garbage=3, deflate=True)
        return dest

    return for_each(files, work, one, "ebooks.zip")


@tool(
    id="text-to-pdf", name="Text to PDF", category="to-pdf", glyph="text",
    description="Plain .txt, code or logs into a clean monospaced PDF.",
    accept=[".txt", ".log", ".py", ".js", ".json", ".xml", ".yaml", ".yml", ".ini", ".conf", ".sh", ".sql", ".c",
            ".cpp", ".java", ".ts", ".css"],
    multiple=True,
    options=[{"name": "size", "label": "Font size", "type": "range", "min": 6, "max": 14, "default": 9}],
)
def text_to_pdf(files, opts, work):
    size = num(opts, "size", 9)

    def one(f):
        text = f.read_text(encoding="utf-8", errors="replace")
        dest = out_path(work, f, ".pdf")
        body = "<pre style='font-family:monospace;font-size:%spt;white-space:pre-wrap'>%s</pre>" % (size, html.escape(text))
        html_to_pdf_story(f"<body>{body}</body>", dest)
        return dest

    return for_each(files, work, one, "text_pdfs.zip")


def chromium(args: list[str], work: Path, timeout: int = 120) -> None:
    run(["chromium", "--headless=new", "--no-sandbox", "--disable-gpu", "--disable-dev-shm-usage",
         "--hide-scrollbars", "--no-first-run", "--disable-extensions", f"--user-data-dir={work / 'chrome'}",
         "--virtual-time-budget=8000", *args], timeout=timeout)


def _fit_pages(src: fitz.Document, paper: str, landscape: bool) -> fitz.Document:
    pr = fitz.paper_rect(paper)
    w, h = (pr.height, pr.width) if landscape else (pr.width, pr.height)
    doc = fitz.open()
    for i in range(src.page_count):
        doc.new_page(width=w, height=h).show_pdf_page(fitz.Rect(0, 0, w, h), src, i)
    return doc


@tool(
    id="url-to-pdf", name="Web page to PDF", category="to-pdf", glyph="globe",
    description="Capture any public web page as a PDF, rendered by Chromium.",
    keywords="website html url link",
    accept=[], min_files=0,
    options=[
        {"name": "url", "label": "URL", "type": "url", "placeholder": "https://", "required": True, "default": ""},
        {"name": "paper", "label": "Page size", "type": "segmented", "default": "a4", "choices": [["a4", "A4"], ["letter", "Letter"]]},
        {"name": "landscape", "label": "Landscape", "type": "checkbox", "default": False},
    ],
)
def url_to_pdf(files, opts, work):
    url = check_public_url(opts.get("url"))
    tmp = work / "chrome.pdf"
    args = ["--no-pdf-header-footer", f"--print-to-pdf={tmp}", url]
    chromium(args, work)
    if not tmp.exists():
        raise ToolError("The page could not be rendered.")
    dest = work / "out" / "webpage.pdf"
    _fit_pages(fitz.open(tmp), opts.get("paper", "a4"), bool(opts.get("landscape"))).save(dest, garbage=3, deflate=True)
    return Result("file", dest)


# --------------------------------------------------------------------------- #
# Editor-backed tools (interactive; handled by app/editor.py)
# --------------------------------------------------------------------------- #

for _id, _name, _glyph, _desc, _kw, _pop in [
    ("edit-pdf", "Edit PDF", "edit",
     "Change existing text in its original font and size, add text, images and shapes.",
     "modify text typo font replace change", True),
    ("sign-pdf", "Sign PDF", "sign", "Draw, type or upload your signature and place it anywhere.", "signature esign", True),
    ("fill-form", "Fill PDF form", "form", "Type into form fields, tick boxes and pick options.", "acroform fields", False),
]:
    tool(id=_id, name=_name, category="edit", glyph=_glyph, description=_desc, accept=PDF, kind="editor",
         popular=_pop, keywords=_kw)(None)
